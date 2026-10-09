"""Generate binary test fixtures (DOCX, text PDF, scanned PDF, corrupt PDF) from sample_resume.txt.

Run from the repository root:  backend/.venv/bin/python scripts/generate_fixtures.py
"""

from __future__ import annotations

from pathlib import Path

import docx
import pymupdf as fitz

ROOT = Path(__file__).resolve().parent.parent
FIX = ROOT / "tests" / "fixtures"


def is_heading(line: str) -> bool:
    return line.isupper() and len(line) < 40


def make_docx(lines: list[str], out: Path) -> None:
    d = docx.Document()
    for line in lines:
        if not line.strip():
            continue
        if is_heading(line):
            d.add_heading(line.title(), level=2)
        elif line.startswith("• "):
            d.add_paragraph(line[2:], style="List Bullet")
        else:
            d.add_paragraph(line)
    d.save(out)


def make_pdf(lines: list[str], out: Path) -> None:
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    y = 50
    for line in lines:
        if y > 800:
            page = doc.new_page(width=595, height=842)
            y = 50
        if not line.strip():
            y += 6
            continue
        size = 15 if y == 50 else (11.5 if is_heading(line) else 9.5)
        font = "hebo" if is_heading(line) or y == 50 else "helv"
        # wrap long lines
        words, cur = line.split(), ""
        for w in words:
            test = (cur + " " + w).strip()
            if fitz.get_text_length(test, fontname=font, fontsize=size) > 495:
                page.insert_text((50, y), cur, fontname=font, fontsize=size)
                y += size + 3
                cur = "  " + w if line.startswith("• ") else w
            else:
                cur = test
        page.insert_text((50, y), cur, fontname=font, fontsize=size)
        y += size + (6 if is_heading(line) else 3)
    doc.save(out)


def make_scanned_pdf(out: Path) -> None:
    """An image-only PDF: rasterise the text PDF so there is no extractable text layer."""
    src = fitz.open(FIX / "sample_resume.pdf")
    pix = src[0].get_pixmap(dpi=50)
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_image(page.rect, stream=pix.tobytes("jpeg", jpg_quality=60))
    doc.save(out)


def main() -> None:
    lines = (FIX / "sample_resume.txt").read_text(encoding="utf-8").splitlines()
    make_docx(lines, FIX / "sample_resume.docx")
    make_pdf(lines, FIX / "sample_resume.pdf")
    make_scanned_pdf(FIX / "scanned_resume.pdf")
    (FIX / "corrupt_resume.pdf").write_bytes(b"%PDF-1.7\n1 0 obj << /Type /Catalog >>\nthis is not a real pdf\n")
    (FIX / "empty_resume.txt").write_text("", encoding="utf-8")
    print("fixtures written to", FIX)


if __name__ == "__main__":
    main()
