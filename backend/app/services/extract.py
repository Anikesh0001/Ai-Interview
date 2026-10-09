"""Local resume text extraction for PDF, DOCX and TXT files.

Nothing is sent anywhere: PDFs are read with PyMuPDF and DOCX with python-docx.
"""

from __future__ import annotations

import io
import logging
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field

logger = logging.getLogger(__name__)

ALLOWED_EXTENSIONS = {"pdf", "docx", "txt", "doc"}
MIN_TEXT_CHARS = 80


class ResumeParseError(Exception):
    """Raised with a message that is safe to show to the user."""

    def __init__(self, user_message: str, status_code: int = 422):
        super().__init__(user_message)
        self.user_message = user_message
        self.status_code = status_code


@dataclass
class ExtractedDocument:
    text: str
    file_type: str
    page_count: int | None = None
    warnings: list[str] = field(default_factory=list)
    signals: dict[str, int | bool] = field(default_factory=dict)


_FILENAME_SAFE = re.compile(r"[^A-Za-z0-9._ -]+")


def sanitize_filename(name: str) -> str:
    """Strip any path components and unsafe characters from a client-supplied filename."""
    base = (name or "").replace("\\", "/").split("/")[-1]
    base = unicodedata.normalize("NFKD", base).encode("ascii", "ignore").decode()
    base = _FILENAME_SAFE.sub("_", base).strip(" .")
    base = re.sub(r"\.{2,}", ".", base)
    return (base or "resume")[:120]


def file_extension(name: str) -> str:
    base = sanitize_filename(name)
    return base.rsplit(".", 1)[-1].lower() if "." in base else ""


_LIGATURES = {"ﬀ": "ff", "ﬁ": "fi", "ﬂ": "fl", "ﬃ": "ffi", "ﬄ": "ffl"}
_BULLETS = "•●▪◦○■□➢➤►▶✓✔◆◇·∙‣⁃–—"


def normalize_text(text: str) -> str:
    for k, v in _LIGATURES.items():
        text = text.replace(k, v)
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ").replace("\t", " ")
    # drop control characters except newlines
    text = "".join(ch for ch in text if ch == "\n" or unicodedata.category(ch)[0] != "C")
    lines = []
    for raw in text.split("\n"):
        line = re.sub(r" {2,}", " ", raw).strip()
        if line and line[0] in _BULLETS and len(line) > 1:
            line = "• " + line[1:].strip()
        lines.append(line)
    out = "\n".join(lines)
    out = re.sub(r"\n{3,}", "\n\n", out)
    return out.strip()


def _extract_pdf(data: bytes) -> ExtractedDocument:
    import pymupdf as fitz

    if not data.startswith(b"%PDF"):
        raise ResumeParseError("This file does not look like a valid PDF. Please re-export it and try again.")
    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:  # corrupted / truncated
        logger.info("pdf open failed: %s", type(exc).__name__)
        raise ResumeParseError("This PDF appears to be corrupted and could not be read.") from exc
    try:
        if doc.needs_pass:
            raise ResumeParseError("This PDF is password-protected. Please upload an unlocked copy.")
        pages: list[str] = []
        image_count = 0
        multi_column_pages = 0
        for page in doc:
            try:
                pages.append(page.get_text("text", sort=True))
                image_count += len(page.get_images(full=False))
                blocks = [b for b in page.get_text("blocks") if b[6] == 0 and len(b[4].strip()) > 20]
                if blocks:
                    width = page.rect.width
                    left = sum(1 for b in blocks if b[2] < width * 0.55)
                    right = sum(1 for b in blocks if b[0] > width * 0.45)
                    if left >= 3 and right >= 3:
                        multi_column_pages += 1
            except Exception as exc:
                logger.info("pdf page extraction failed: %s", type(exc).__name__)
                raise ResumeParseError("This PDF appears to be corrupted and could not be read.") from exc
        page_count = doc.page_count
        if page_count == 0:
            raise ResumeParseError("This PDF appears to be corrupted and could not be read.")
    finally:
        doc.close()

    text = normalize_text("\n".join(pages))
    if len(re.sub(r"\s", "", text)) < MIN_TEXT_CHARS:
        if page_count and (image_count > 0 or not text):
            raise ResumeParseError(
                "Your PDF appears to be scanned/image-based. Please upload a text-based PDF or DOCX."
            )
        raise ResumeParseError("The uploaded resume appears to be empty. Please check the file.")
    warnings = []
    if multi_column_pages:
        warnings.append("Multi-column layout detected; some ATS parsers read columns out of order.")
    if image_count:
        warnings.append("Images or graphics detected; ATS parsers ignore text inside images.")
    return ExtractedDocument(
        text=text,
        file_type="pdf",
        page_count=page_count,
        warnings=warnings,
        signals={"images": image_count, "multi_column_pages": multi_column_pages, "tables": 0},
    )


def _extract_docx(data: bytes) -> ExtractedDocument:
    import docx  # python-docx

    if not data.startswith(b"PK"):
        raise ResumeParseError("This file does not look like a valid DOCX document.")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            if "word/document.xml" not in zf.namelist():
                raise ResumeParseError("This file does not look like a valid DOCX document.")
        document = docx.Document(io.BytesIO(data))
    except ResumeParseError:
        raise
    except Exception as exc:
        logger.info("docx open failed: %s", type(exc).__name__)
        raise ResumeParseError("This DOCX file appears to be corrupted and could not be read.") from exc

    parts: list[str] = []
    for para in document.paragraphs:
        txt = para.text
        style = (para.style.name or "").lower() if para.style is not None else ""
        if "list" in style and txt.strip() and txt.strip()[0] not in _BULLETS:
            txt = "• " + txt
        parts.append(txt)
    table_count = len(document.tables)
    for table in document.tables:
        for row in table.rows:
            seen: set[str] = set()
            cells = []
            for cell in row.cells:
                t = cell.text.strip()
                if t and t not in seen:
                    seen.add(t)
                    cells.append(t)
            if cells:
                parts.append(" | ".join(cells))
    header_text = 0
    for section in document.sections:
        for p in section.header.paragraphs:
            header_text += len(p.text.strip())
    image_count = len(document.inline_shapes)

    text = normalize_text("\n".join(parts))
    if len(re.sub(r"\s", "", text)) < MIN_TEXT_CHARS:
        raise ResumeParseError("The uploaded resume appears to be empty. Please check the file.")
    warnings = []
    if table_count:
        warnings.append("Tables detected; many ATS parsers flatten or skip table content.")
    if header_text:
        warnings.append("Text found in the page header; some ATS parsers ignore headers and footers.")
    if image_count:
        warnings.append("Images detected; ATS parsers ignore text inside images.")
    return ExtractedDocument(
        text=text,
        file_type="docx",
        page_count=None,
        warnings=warnings,
        signals={"tables": table_count, "images": image_count, "header_text": header_text},
    )


def _extract_txt(data: bytes) -> ExtractedDocument:
    if b"\x00" in data[:4096] and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
        raise ResumeParseError("This text file contains binary data. Please upload a plain-text resume.")
    text = None
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            if enc == "utf-16" and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
                continue
            break
        except (UnicodeDecodeError, UnicodeError):
            continue
    if text is None:
        raise ResumeParseError("Could not decode this text file. Please save it as UTF-8.")
    text = normalize_text(text)
    if len(re.sub(r"\s", "", text)) < MIN_TEXT_CHARS:
        raise ResumeParseError("The uploaded resume appears to be empty. Please check the file.")
    return ExtractedDocument(text=text, file_type="txt", signals={"tables": 0, "images": 0})


def extract_text(filename: str, data: bytes, max_bytes: int) -> ExtractedDocument:
    """Validate and extract text from an uploaded resume."""
    ext = file_extension(filename)
    if ext not in ALLOWED_EXTENSIONS:
        raise ResumeParseError("Unsupported file type. Please upload a PDF, DOCX or TXT resume.", 415)
    if not data:
        raise ResumeParseError("The uploaded file is empty.")
    if len(data) > max_bytes:
        raise ResumeParseError(f"The file is too large. The maximum size is {max_bytes // (1024 * 1024)} MB.", 413)
    if ext == "doc":
        raise ResumeParseError(
            "Legacy .doc files can't be parsed reliably offline. Please save the resume as DOCX or PDF.", 415
        )
    if ext == "pdf":
        return _extract_pdf(data)
    if ext == "docx":
        return _extract_docx(data)
    return _extract_txt(data)
